import os

import openpyxl

from backend import config
from backend.grader import validate_and_grade


def test_grader():
    wb = openpyxl.Workbook()
    ws_meta = wb.create_sheet("Metadata")
    ws_meta.append(["Question_Number", "Chapter", "Correct_Answer"])
    for i in range(1, 61):
        ws_meta.append([i, "Algebra" if i <= 30 else "Geometry", "A"])

    ws_entry = wb.create_sheet("Entry")
    ws_entry.append(["Student ID", "Name", "Date of Birth"] + [f"Q{i}" for i in range(1, 61)])

    # Student 1: 60 correct — new format: "A - Correct"
    ws_entry.append(["S1", "Alice", "2010-01-01"] + ["A - Correct"] * 60)

    # Student 2: 10 correct, 50 incorrect. Raw = -40. Floor = 0.
    ws_entry.append(["S2", "Bob", "2010-02-02"] + ["A - Correct"] * 10 + ["Not A - Incorrect"] * 50)

    # Student 3: Blank Q5
    ws_entry.append(["S3", "Charlie", "2010-03-03"] + ["A - Correct"] * 4 + [""] + ["A - Correct"] * 55)

    temp_file = "test_grader.xlsx"
    wb.save(temp_file)

    res = validate_and_grade(temp_file)
    assert not res["success"]
    assert any("Row 4: Q5 is blank" in err for err in res["errors"])

    # Fix the blank error
    ws_entry.cell(row=4, column=8, value="Unattempted")  # Q5 is column H (8th col)
    wb.save(temp_file)

    res = validate_and_grade(temp_file)
    assert res["success"]

    students = res["students"]
    assert len(students) == 3

    alice = next(s for s in students if s["name"] == "Alice")
    assert alice["final_score"] == 60
    assert alice["velocity_index"] == round(60 / config.EXAM_DURATION_MINUTES, 2)
    assert alice["rank"] == 1
    assert alice["topic_breakdown"]["Algebra"]["percentage"] == 100.0

    bob = next(s for s in students if s["name"] == "Bob")
    assert bob["final_score"] == 0  # Flooring verification
    assert bob["velocity_index"] == round(60 / config.EXAM_DURATION_MINUTES, 2)

    charlie = next(s for s in students if s["name"] == "Charlie")
    assert charlie["final_score"] == 59
    assert charlie["velocity_index"] == round(59 / config.EXAM_DURATION_MINUTES, 2)

    if os.path.exists(temp_file):
        os.remove(temp_file)


def test_blank_row_in_middle_does_not_drop_later_students(tmp_path):
    wb = openpyxl.Workbook()
    ws_meta = wb.create_sheet("Metadata")
    ws_meta.append(["Question_Number", "Chapter", "Correct_Answer"])
    for i in range(1, 61):
        ws_meta.append([i, "Algebra", "A"])

    ws_entry = wb.create_sheet("Entry")
    ws_entry.append(["Student ID", "Name", "Date of Birth"] + [f"Q{i}" for i in range(1, 61)])
    ws_entry.append(["S1", "Alice", "2010-01-01"] + ["A - Correct"] * 60)
    ws_entry.append([None] * 63)  # stray blank row in the middle
    ws_entry.append(["S2", "Bob", "2010-02-02"] + ["A - Correct"] * 60)
    ws_entry.append(["S3", "Charlie", "2010-03-03"] + ["A - Correct"] * 60)

    temp_file = str(tmp_path / "blank_row.xlsx")
    wb.save(temp_file)

    res = validate_and_grade(temp_file)
    assert res["success"], res.get("errors")
    names = {s["name"] for s in res["students"]}
    assert names == {"Alice", "Bob", "Charlie"}
    assert res["rows_read"] == 3


def test_missing_chapter_becomes_unknown(tmp_path):
    wb = openpyxl.Workbook()
    ws_meta = wb.create_sheet("Metadata")
    ws_meta.append(["Question_Number", "Chapter", "Correct_Answer"])
    for i in range(1, 61):
        # Leave the chapter blank for the first 10 questions.
        chapter = None if i <= 10 else "Geometry"
        ws_meta.append([i, chapter, "A"])

    ws_entry = wb.create_sheet("Entry")
    ws_entry.append(["Student ID", "Name", "Date of Birth"] + [f"Q{i}" for i in range(1, 61)])
    ws_entry.append(["S1", "Alice", "2010-01-01"] + ["A - Correct"] * 60)

    temp_file = str(tmp_path / "missing_chapter.xlsx")
    wb.save(temp_file)

    res = validate_and_grade(temp_file)
    assert res["success"], res.get("errors")
    alice = res["students"][0]
    assert "Unknown" in alice["topic_breakdown"]
    assert "None" not in alice["topic_breakdown"]
    assert alice["topic_breakdown"]["Unknown"]["total"] == 10
    assert alice["topic_breakdown"]["Geometry"]["total"] == 50


def test_negative_raw_score_clamps_to_zero_final(tmp_path):
    wb = openpyxl.Workbook()
    ws_meta = wb.create_sheet("Metadata")
    ws_meta.append(["Question_Number", "Chapter", "Correct_Answer"])
    for i in range(1, 61):
        ws_meta.append([i, "Algebra", "A"])

    ws_entry = wb.create_sheet("Entry")
    ws_entry.append(["Student ID", "Name", "Date of Birth"] + [f"Q{i}" for i in range(1, 61)])
    # 5 correct, 20 incorrect, 35 unattempted -> raw_score = -15
    responses = ["A - Correct"] * 5 + ["Not A - Incorrect"] * 20 + ["Unattempted"] * 35
    ws_entry.append(["S1", "Alice", "2010-01-01"] + responses)

    temp_file = str(tmp_path / "negative_score.xlsx")
    wb.save(temp_file)

    res = validate_and_grade(temp_file)
    assert res["success"], res.get("errors")
    alice = res["students"][0]
    assert alice["raw_score"] == -15
    assert alice["final_score"] == 0
