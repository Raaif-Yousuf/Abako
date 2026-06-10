import os
import openpyxl
from backend.grader import validate_and_grade

def test_corrupted_excel():
    filename = "corrupted.xlsx"
    with open(filename, "w") as f:
        f.write("This is completely invalid data.")

    res = validate_and_grade(filename)
    assert res["success"] == False
    assert any("Failed to open Excel file" in e for e in res["errors"])
    os.remove(filename)

def test_schema_mismatch():
    filename = "schema_mismatch.xlsx"
    wb = openpyxl.Workbook()
    wb.create_sheet("Metadata")
    ws_entry = wb.create_sheet("Entry")

    # Intentionally missing columns after Q2
    ws_entry.append(["Student ID", "Name", "Date of Birth", "Q1", "Q2"])
    ws_entry.append(["S1", "Alice", "2010", "A - Correct", "Not A - Incorrect"])
    wb.save(filename)

    res = validate_and_grade(filename)
    assert res["success"] == False
    assert any("Missing columns" in e for e in res["errors"])
    os.remove(filename)

def test_extreme_data():
    filename_1 = "extreme_1.xlsx"
    wb = openpyxl.Workbook()
    ws_meta = wb.create_sheet("Metadata")
    for i in range(1, 61): ws_meta.append([i, "Algebra", "A"])
    ws_entry = wb.create_sheet("Entry")
    ws_entry.append(["Student ID", "Name", "Date of Birth"] + [f"Q{i}" for i in range(1, 61)])

    # 1 student — new format
    ws_entry.append(["S1", "Alice", "2010-01-01"] + ["A - Correct"] * 60)
    wb.save(filename_1)
    res1 = validate_and_grade(filename_1)
    assert res1["success"] == True
    assert res1["students"][0]["percentile"] == 100.0
    os.remove(filename_1)

    # 1000 students
    filename_1000 = "extreme_1000.xlsx"
    wb2 = openpyxl.Workbook()
    wb2.create_sheet("Metadata")
    for i in range(1, 61): wb2["Metadata"].append([i, "Algebra", "A"])
    ws2 = wb2.create_sheet("Entry")
    ws2.append(["Student ID", "Name", "Date of Birth"] + [f"Q{i}" for i in range(1, 61)])

    for i in range(1000):
        ws2.append([f"S{i}", f"Name{i}", "2010"] + ["A - Correct"] * 60)
    wb2.save(filename_1000)

    res2 = validate_and_grade(filename_1000)
    assert res2["success"] == True
    assert len(res2["students"]) == 1000
    os.remove(filename_1000)
