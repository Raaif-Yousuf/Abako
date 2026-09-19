"""Pins down the percentile-rank and competition-ranking fix in backend.grader.

Percentile definition used: the percentage of the cohort scoring strictly
below a given student's final_score, rounded to one decimal. Tied students
therefore always receive the same percentile. A cohort of exactly one
student always gets 100.0.

Rank uses competition ranking ("1224"): tied students share the better rank
and the next distinct score skips the tied places.
"""

import openpyxl

from backend.grader import validate_and_grade


def _build_workbook(path, scores):
    """Build a gradeable workbook where student i's final_score == scores[i].

    Each student answers `score` questions correctly and leaves the rest
    unattempted, so raw_score == final_score == score (no negative clamping
    involved).
    """
    wb = openpyxl.Workbook()
    ws_meta = wb.create_sheet("Metadata")
    ws_meta.append(["Question_Number", "Chapter", "Correct_Answer"])
    for i in range(1, 61):
        ws_meta.append([i, "General", "A"])

    ws_entry = wb.create_sheet("Entry")
    ws_entry.append(["Student ID", "Name", "Date of Birth"] + [f"Q{i}" for i in range(1, 61)])

    for idx, score in enumerate(scores, start=1):
        responses = ["A - Correct"] * score + ["Unattempted"] * (60 - score)
        ws_entry.append([f"S{idx}", f"Student{idx}", "2010-01-01"] + responses)

    wb.save(path)


def _grade(tmp_path, scores, filename="grade.xlsx"):
    path = tmp_path / filename
    _build_workbook(str(path), scores)
    result = validate_and_grade(str(path))
    assert result["success"], result.get("errors")
    return {s["name"]: s for s in result["students"]}


def test_all_tied_cohort_gets_same_percentile(tmp_path):
    students = _grade(tmp_path, [30, 30, 30, 30, 30])
    percentiles = {s["percentile"] for s in students.values()}
    ranks = {s["rank"] for s in students.values()}
    assert percentiles == {0.0}
    assert ranks == {1}


def test_cohort_of_one_is_100th_percentile(tmp_path):
    students = _grade(tmp_path, [17])
    student = students["Student1"]
    assert student["percentile"] == 100.0
    assert student["rank"] == 1


def test_cohort_of_two(tmp_path):
    students = _grade(tmp_path, [40, 20])
    top = students["Student1"]
    bottom = students["Student2"]
    assert top["rank"] == 1
    assert top["percentile"] == 50.0
    assert bottom["rank"] == 2
    assert bottom["percentile"] == 0.0


def test_normal_spread_exact_values(tmp_path):
    # Scores (desc): 60, 45, 45, 30, 15, 15, 15, 0 -- n=8.
    scores = [60, 45, 45, 30, 15, 15, 15, 0]
    students = _grade(tmp_path, scores)

    expected = {
        "Student1": (60, 1, 87.5),
        "Student2": (45, 2, 62.5),
        "Student3": (45, 2, 62.5),
        "Student4": (30, 4, 50.0),
        "Student5": (15, 5, 12.5),
        "Student6": (15, 5, 12.5),
        "Student7": (15, 5, 12.5),
        "Student8": (0, 8, 0.0),
    }
    for name, (score, rank, percentile) in expected.items():
        s = students[name]
        assert s["final_score"] == score
        assert s["rank"] == rank
        assert s["percentile"] == percentile


def test_top_scorer_highest_bottom_scorer_zero(tmp_path):
    scores = [60, 45, 45, 30, 15, 15, 15, 0]
    students = _grade(tmp_path, scores)
    top_percentile = students["Student1"]["percentile"]
    bottom_percentile = students["Student8"]["percentile"]
    all_percentiles = [s["percentile"] for s in students.values()]
    assert top_percentile == max(all_percentiles)
    assert bottom_percentile == 0.0


def test_competition_ranking_two_tie_groups(tmp_path):
    # Scores (desc): 50, 50, 50, 30, 30, 10 -- two separate groups of ties
    # plus a unique bottom score.
    scores = [50, 50, 50, 30, 30, 10]
    students = _grade(tmp_path, scores)

    for name in ("Student1", "Student2", "Student3"):
        assert students[name]["rank"] == 1
    for name in ("Student4", "Student5"):
        assert students[name]["rank"] == 4
    assert students["Student6"]["rank"] == 6
