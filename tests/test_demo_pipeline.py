"""End-to-end check on the built-in demo.

The demo is the first thing anyone clicks, so it is worth a test of its own:
it walks the whole pipeline (select questions -> question paper -> answer key
-> entry sheet -> fill it -> grade -> reports) and is the only place where a
mismatch between the entry sheet's dropdown values and the grader's accepted
values shows up.
"""

import os
import shutil

import openpyxl
import pytest

from backend.demo_manager import run_demo
from backend.grader import validate_and_grade

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


@pytest.fixture
def demo(tmp_path, monkeypatch):
    """Run the demo inside a throwaway directory with a copy of the resources."""
    resources = tmp_path / "resources"
    resources.mkdir()
    for name in ("question_bank.xlsx", "abako_logo.png"):
        src = os.path.join(REPO_ROOT, "resources", name)
        if os.path.exists(src):
            shutil.copy2(src, resources / name)

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        "backend.demo_manager.QUESTION_BANK_PATH", str(resources / "question_bank.xlsx")
    )
    return run_demo()


def test_demo_runs_end_to_end(demo):
    assert demo["success"], demo.get("errors")


def test_demo_entry_sheet_grades_cleanly(demo):
    """Regression: the demo used to write bare "Correct" / "Incorrect", which
    the grader rejects, so every demo run failed validation."""
    assert demo["success"], demo.get("errors")
    sheet = os.path.join(demo["output_dir"], "Data_Entry_Sheet.xlsx")
    result = validate_and_grade(sheet)
    assert result["success"], result.get("errors", [])[:5]
    assert len(result["students"]) == 20


def test_demo_entry_values_match_the_dropdowns(demo):
    """Every filled cell must be one of the three options its dropdown offers."""
    assert demo["success"], demo.get("errors")
    sheet = os.path.join(demo["output_dir"], "Data_Entry_Sheet.xlsx")
    wb = openpyxl.load_workbook(sheet, data_only=True)
    answers = [row[2] for row in wb["Metadata"].iter_rows(min_row=3, values_only=True)]

    rows = list(wb["Entry"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 20
    for row in rows:
        for q_idx, value in enumerate(row[3:63]):
            allowed = {
                f"{answers[q_idx]} - Correct",
                f"Not {answers[q_idx]} - Incorrect",
                "Unattempted",
            }
            assert value in allowed, f"Q{q_idx + 1} got {value!r}"


def test_demo_produces_one_report_per_student(demo):
    assert demo["success"], demo.get("errors")
    produced = []
    indiv_dir = os.path.join(demo["output_dir"], "Individual_Reports")
    for _root, _dirs, files in os.walk(demo["output_dir"]):
        produced += [f for f in files if f.endswith(".pdf")]
    assert "Question_Paper.pdf" in produced
    assert "Answer_Key.pdf" in produced
    assert "School_Report.pdf" in produced
    assert len(os.listdir(indiv_dir)) == 20
