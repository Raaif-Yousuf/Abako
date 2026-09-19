import os

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

# Matches the brand palette documented in backend/pdf_theme.py (charcoal and
# the pale hairline/panel tints). openpyxl fills are plain "RRGGBB" strings,
# not reportlab Color objects, so the values are restated here rather than
# imported.
_CHARCOAL = "363434"
_HEADER_TEXT = "FFFFFF"
_UNFILLED_FILL = "FFF3CD"


def generate_data_entry_sheet(filename, questions, school_name='', campus_name='', grade=''):
    os.makedirs(os.path.dirname(filename) or '.', exist_ok=True)

    wb = Workbook()

    # --- Tab 1: Instructions ---
    ws_instr = wb.active
    ws_instr.title = "Instructions"
    ws_instr.append(["Data Entry Instructions"])
    ws_instr.append([])
    ws_instr.append(["1. Do not rename any columns or move the sheets."])
    ws_instr.append(["2. Fill in Student ID, Name, and Date of Birth for every student."])
    ws_instr.append(["3. For Q1 to Q60, use the dropdown in each cell. Each question's dropdown "
                      "shows its correct answer as one of the three options."])
    ws_instr.append(["4. A pale highlight marks a cell that still needs an answer once a row has "
                      "a Student ID - fill it in or select Unattempted."])
    ws_instr.append(["5. Do not modify, delete, or unhide the Metadata sheet - doing so breaks the grader."])

    ws_instr.column_dimensions['A'].width = 100
    ws_instr['A1'].font = Font(bold=True, size=14, color=_CHARCOAL)
    for row in range(3, 6 + 2):
        cell = ws_instr.cell(row=row, column=1)
        if cell.value:
            cell.font = Font(size=11)
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    # --- Tab 2: Entry ---
    ws_entry = wb.create_sheet("Entry")

    q_headers = [f"Q{i+1}" for i in range(60)]
    headers = ["Student ID", "Name", "Date of Birth"] + q_headers
    ws_entry.append(headers)

    header_fill = PatternFill(start_color=_CHARCOAL, end_color=_CHARCOAL, fill_type="solid")
    header_font = Font(color=_HEADER_TEXT, bold=True)
    header_border = Border(bottom=Side(style="thin", color=_CHARCOAL))
    for cell in ws_entry[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.border = header_border
        cell.alignment = Alignment(horizontal='center', vertical='center')

    ws_entry.freeze_panes = "D2"
    ws_entry.row_dimensions[1].height = 20

    # Per-question dynamic dropdowns: "X - Correct, Not X - Incorrect, Unattempted"
    for i, q in enumerate(questions):
        col_idx = i + 4  # Q1 -> column D (4)
        col_letter = ws_entry.cell(row=1, column=col_idx).column_letter
        answer = str(q.get('Correct_Answer', ''))
        formula = f'"{answer} - Correct,Not {answer} - Incorrect,Unattempted"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        dv.error = 'Your entry is not in the list'
        dv.errorTitle = 'Invalid Entry'
        dv.prompt = 'Select from the dropdown'
        dv.promptTitle = 'Validation'
        ws_entry.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}1000")

    ws_entry.column_dimensions['A'].width = 14
    ws_entry.column_dimensions['B'].width = 25
    ws_entry.column_dimensions['C'].width = 15
    last_col_letter = ws_entry.cell(row=1, column=3 + 60).column_letter
    for col_idx in range(4, 64):
        col_letter = ws_entry.cell(row=1, column=col_idx).column_letter
        ws_entry.column_dimensions[col_letter].width = 15

    # Make an unfilled cell obvious: once a row has a Student ID, any Name,
    # Date of Birth or question cell still blank in that row is highlighted.
    unfilled_fill = PatternFill(start_color=_UNFILLED_FILL, end_color=_UNFILLED_FILL,
                                 fill_type="solid")
    ws_entry.conditional_formatting.add(
        f"B2:{last_col_letter}1000",
        FormulaRule(formula=["AND($A2<>\"\",B2=\"\")"], fill=unfilled_fill),
    )

    # --- Tab 3: Metadata (Hidden) ---
    ws_meta = wb.create_sheet("Metadata")
    # Row 1: exam info marker (read by /read-metadata endpoint)
    ws_meta.append(["EXAM_META", school_name, campus_name, grade])
    # Row 2: question data header
    ws_meta.append(["Question_Number", "Chapter", "Correct_Answer"])
    for i, q in enumerate(questions):
        chapter = q.get('Chapter', 'Unknown')
        ans = q.get('Correct_Answer', '')
        ws_meta.append([i + 1, chapter, ans])

    ws_meta.sheet_state = 'hidden'

    wb.save(filename)
