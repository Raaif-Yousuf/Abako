import sys
import os
import json
import socket
import shutil
import logging

# Ensure project root is in path when invoked as `python backend/main.py`
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from flask import Flask, request, jsonify
from flask_cors import CORS

# Silence Werkzeug so only our startup line goes to stdout
logging.getLogger("werkzeug").disabled = True
logging.getLogger("werkzeug").handlers = []

app = Flask(__name__)
CORS(app)

from backend.config import QUESTION_BANK_PATH, OUTPUT_BASE_PATH
from backend.generator import generate_from_excel
from backend.pdf_manager import generate_question_paper, generate_answer_key
from backend.excel_manager import generate_data_entry_sheet
from backend.grader import validate_and_grade
from backend.reporter import generate_all_reports
from backend.demo_manager import run_demo


def _find_free_port():
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@app.route("/ping", methods=["POST"])
def ping():
    return jsonify({"status": "ok"})


@app.route("/check-bank", methods=["POST"])
def check_bank():
    exists = os.path.exists(QUESTION_BANK_PATH)
    filename = os.path.basename(QUESTION_BANK_PATH) if exists else None
    return jsonify({"exists": exists, "filename": filename})


@app.route("/save-bank", methods=["POST"])
def save_bank():
    data = request.get_json() or {}
    source = data.get("source_path", "")
    if not os.path.exists(source):
        return jsonify({"status": "error", "message": "Source file not found"}), 400
    os.makedirs(os.path.dirname(QUESTION_BANK_PATH) or ".", exist_ok=True)
    shutil.copy2(source, QUESTION_BANK_PATH)
    return jsonify({"status": "ok"})


@app.route("/get-chapters", methods=["POST"])
def get_chapters():
    if not os.path.exists(QUESTION_BANK_PATH):
        return jsonify({"chapters": []})
    import pandas as pd
    xl = pd.ExcelFile(QUESTION_BANK_PATH)
    chapters = [s for s in xl.sheet_names if s != "Filler"]
    return jsonify({"chapters": chapters})


@app.route("/generate", methods=["POST"])
def generate():
    data = request.get_json() or {}
    school_name = data.get("school_name", "")
    campus_name = data.get("campus_name", "")
    grade = data.get("grade", "")
    exam_date = data.get("exam_date", "")
    chapter_allocations = data.get("chapter_allocations", {})
    shuffle = data.get("shuffle", True)
    selected_chapters = list(chapter_allocations.keys())

    school_info = {
        "School Name": school_name,
        "Campus": campus_name,
        "Grade": grade,
        "Exam Date": exam_date,
    }

    safe_school = "".join(c if c.isalnum() else "_" for c in school_name)
    safe_campus = "".join(c if c.isalnum() else "_" for c in campus_name)
    safe_grade = "".join(c if c.isalnum() else "_" for c in grade)
    output_dir = os.path.join(OUTPUT_BASE_PATH, f"{safe_school}_{safe_campus}_{safe_grade}")
    os.makedirs(output_dir, exist_ok=True)

    questions = generate_from_excel(QUESTION_BANK_PATH, selected_chapters, chapter_allocations=chapter_allocations, shuffle=shuffle)

    q_pdf = os.path.join(output_dir, "Question_Paper.pdf")
    a_pdf = os.path.join(output_dir, "Answer_Key.pdf")
    e_xlsx = os.path.join(output_dir, "Data_Entry_Sheet.xlsx")

    generate_question_paper(q_pdf, school_info, questions)
    generate_answer_key(a_pdf, questions, school_name=school_name, campus_name=campus_name, grade=grade, exam_date=exam_date)
    generate_data_entry_sheet(e_xlsx, questions, school_name=school_name, campus_name=campus_name, grade=grade)

    return jsonify({"output_dir": output_dir, "files": [q_pdf, a_pdf, e_xlsx]})


@app.route("/read-metadata", methods=["POST"])
def read_metadata():
    data = request.get_json() or {}
    file_path = data.get("path", "")
    if not os.path.exists(file_path):
        return jsonify({"success": False, "error": "File not found"}), 400
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, data_only=True)
        if "Metadata" not in wb.sheetnames:
            return jsonify({"success": False, "error": "No Metadata sheet"})
        ws = wb["Metadata"]
        marker = ws.cell(row=1, column=1).value
        if marker != "EXAM_META":
            return jsonify({"success": False, "error": "No exam info in metadata"})
        school_name = str(ws.cell(row=1, column=2).value or "")
        campus_name = str(ws.cell(row=1, column=3).value or "")
        grade = str(ws.cell(row=1, column=4).value or "")
        return jsonify({"success": True, "school_name": school_name, "campus_name": campus_name, "grade": grade})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500


@app.route("/run-demo", methods=["POST"])
def run_demo_endpoint():
    result = run_demo()
    if result.get("success"):
        return jsonify({"status": "ok", "output_dir": result["output_dir"]})
    return jsonify({"status": "error", "errors": result.get("errors", ["Demo failed"])}), 500


@app.route("/grade", methods=["POST"])
def grade():
    data = request.get_json() or {}
    entry_sheet_path = data.get("entry_sheet_path", "")
    output_folder = data.get("output_folder", OUTPUT_BASE_PATH)
    school_info = {
        "School Name": data.get("school_name", ""),
        "Campus": data.get("campus_name", ""),
        "Grade": data.get("grade", ""),
    }

    if not os.path.exists(entry_sheet_path):
        return jsonify({"success": False, "errors": ["Entry sheet not found"]}), 400

    result = validate_and_grade(entry_sheet_path)
    if not result.get("success"):
        return jsonify(result), 400

    generate_all_reports(result, school_info, base_dir=output_folder)

    return jsonify({"output_dir": output_folder})


class _NullWriter:
    def write(self, *a, **kw): pass
    def flush(self, *a, **kw): pass

if __name__ == "__main__":
    port = _find_free_port()
    sys.stdout.write(json.dumps({"status": "ready", "port": port}) + "\n")
    sys.stdout.flush()
    # Swallow all subsequent stdout (Werkzeug banners go via click.echo → stdout)
    sys.stdout = _NullWriter()
    app.run(host="127.0.0.1", port=port, debug=False, use_reloader=False)
