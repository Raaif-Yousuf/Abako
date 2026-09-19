from collections import Counter

import openpyxl

from backend import config

MAX_REPORTED_ERRORS = 25


def validate_and_grade(filepath):
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
    except Exception as e:
        return {"success": False, "errors": [f"Failed to open Excel file (might be corrupted): {str(e)}"]}

    if "Entry" not in wb.sheetnames or "Metadata" not in wb.sheetnames:
        return {"success": False, "errors": ["Missing required sheets 'Entry' or 'Metadata'."]}

    ws_entry = wb["Entry"]
    ws_meta = wb["Metadata"]

    # Read metadata
    # Row 1 may be EXAM_META (exam info); row 2 is the question header.
    # Use try/except to skip any non-integer rows (EXAM_META, header row).
    topic_mapping = {}
    for row in ws_meta.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        try:
            q_num = int(row[0])
        except (ValueError, TypeError):
            continue
        raw_chapter = row[1] if len(row) > 1 else None
        chapter = "Unknown" if raw_chapter is None or str(raw_chapter).strip() == "" else str(raw_chapter)
        topic_mapping[q_num] = chapter

    # Read Entry
    errors = []
    students_data = []
    rows_read = 0

    for row_idx, row in enumerate(ws_entry.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            # A stray blank row shouldn't truncate the rest of the class list.
            continue
        rows_read += 1

        student_id, name = row[0], row[1]
        if not student_id:
            errors.append(f"Row {row_idx}: Missing Student ID")
        if not name:
            errors.append(f"Row {row_idx}: Missing Name")

        q_responses = row[3:63]
        if len(q_responses) < 60:
            errors.append(f"Row {row_idx}: Missing columns for questions")
            continue

        for q_idx in range(60):
            val = q_responses[q_idx]
            if val is None or val == '':
                errors.append(f"Row {row_idx}: Q{q_idx+1} is blank")
            elif not (val == 'Unattempted' or str(val).endswith('- Correct') or str(val).endswith('- Incorrect')):
                errors.append(f"Row {row_idx}: Q{q_idx+1} has invalid value '{val}'")

        if not errors:
            students_data.append({
                "student_id": student_id,
                "name": name,
                "responses": q_responses
            })

    if errors:
        if len(errors) > MAX_REPORTED_ERRORS:
            suppressed = len(errors) - MAX_REPORTED_ERRORS
            errors = errors[:MAX_REPORTED_ERRORS]
            errors.append(f"...and {suppressed} more error(s) suppressed.")
        return {"success": False, "errors": errors}

    # Grade
    graded_students = []
    for s in students_data:
        correct = 0
        incorrect = 0
        topic_stats = {ch: {"correct": 0, "incorrect": 0, "unattempted": 0, "total": 0} for ch in set(topic_mapping.values())}

        for q_idx in range(60):
            q_num = q_idx + 1
            ans = s["responses"][q_idx]
            chapter = topic_mapping.get(q_num, "Unknown")

            if chapter not in topic_stats:
                topic_stats[chapter] = {"correct": 0, "incorrect": 0, "unattempted": 0, "total": 0}

            topic_stats[chapter]["total"] += 1
            if str(ans).endswith('- Correct'):
                correct += 1
                topic_stats[chapter]["correct"] += 1
            elif str(ans).endswith('- Incorrect'):
                incorrect += 1
                topic_stats[chapter]["incorrect"] += 1
            else:  # Unattempted
                topic_stats[chapter]["unattempted"] += 1

        raw_score = correct - incorrect
        final_score = max(0, raw_score)

        attempts = correct + incorrect
        velocity_index = round(attempts / config.EXAM_DURATION_MINUTES, 2)

        topic_breakdown = {}
        for chapter, stats in topic_stats.items():
            if stats["total"] > 0:
                topic_breakdown[chapter] = {
                    "percentage": round((stats["correct"] / stats["total"]) * 100, 1),
                    "correct": stats["correct"],
                    "incorrect": stats["incorrect"],
                    "unattempted": stats["unattempted"],
                    "total": stats["total"]
                }

        graded_students.append({
            "student_id": s["student_id"],
            "name": s["name"],
            "raw_score": raw_score,
            "final_score": final_score,
            "velocity_index": velocity_index,
            "topic_breakdown": topic_breakdown
        })

    # Rank and Percentile
    graded_students.sort(key=lambda x: x["final_score"], reverse=True)
    n = len(graded_students)
    counts = Counter(s["final_score"] for s in graded_students)

    # Competition ranking ("1224"): tied students share the better rank and the
    # next distinct score skips the tied places (e.g. 1, 2, 2, 4).
    rank_for_score = {}
    higher_count = 0
    for score in sorted(counts.keys(), reverse=True):
        rank_for_score[score] = higher_count + 1
        higher_count += counts[score]

    # Percentile rank = percentage of the cohort scoring strictly below this
    # student's score, rounded to one decimal. Tied students always get the
    # same percentile, and the top scorer approaches (but need not equal) 100.
    if n == 1:
        percentile_for_score = {graded_students[0]["final_score"]: 100.0}
    else:
        percentile_for_score = {}
        lower_count = 0
        for score in sorted(counts.keys()):
            percentile_for_score[score] = round((lower_count / n) * 100, 1)
            lower_count += counts[score]

    for s in graded_students:
        s["rank"] = rank_for_score[s["final_score"]]
        s["percentile"] = percentile_for_score[s["final_score"]]

    return {
        "success": True,
        "students": graded_students,
        "tested_topics": list(set(topic_mapping.values())),
        "rows_read": rows_read
    }
