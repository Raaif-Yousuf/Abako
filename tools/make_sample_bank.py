import os

import openpyxl

CHAPTERS = ["Algebra", "Geometry", "Arithmetic"]
QUESTIONS_PER_CHAPTER = 20
FILLER_QUESTIONS = 15

def create_test_bank():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for chapter in CHAPTERS:
        ws = wb.create_sheet(chapter)
        ws.append(["Question_ID", "Question_Text", "Correct_Answer"])
        for j in range(1, QUESTIONS_PER_CHAPTER + 1):
            q_id = f"{chapter[:3].upper()}_{j:02d}"
            q_text = f"[{chapter}] If x = {j}, what is {j}x + {j}?"
            answer = str(j * j + j)
            ws.append([q_id, q_text, answer])

    ws_filler = wb.create_sheet("Filler")
    ws_filler.append(["Question_ID", "Question_Text", "Correct_Answer"])
    for j in range(1, FILLER_QUESTIONS + 1):
        q_id = f"FIL_{j:02d}"
        q_text = f"[Filler] What is {j} × {j + 1}?"
        answer = str(j * (j + 1))
        ws_filler.append([q_id, q_text, answer])

    output_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "resources", "question_bank.xlsx")
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"Test bank created: {output_path}")
    print(f"  Chapters: {CHAPTERS} ({QUESTIONS_PER_CHAPTER} questions each)")
    print(f"  Filler: {FILLER_QUESTIONS} questions")

if __name__ == "__main__":
    create_test_bank()
