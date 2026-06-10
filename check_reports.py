import json
import os
import openpyxl
from backend.grader import validate_and_grade
from backend.reporter import generate_all_reports

def test_reports():
    temp_file = "test_grader_reports.xlsx"
    wb = openpyxl.Workbook()
    ws_meta = wb.create_sheet("Metadata")
    ws_meta.append(["Question_Number", "Chapter", "Correct_Answer"])
    for i in range(1, 61):
        ws_meta.append([i, "Algebra" if i <= 30 else "Geometry", "A"])
        
    ws_entry = wb.create_sheet("Entry")
    ws_entry.append(["Student ID", "Name", "Date of Birth"] + [f"Q{i}" for i in range(1, 61)])
    
    ws_entry.append(["S1", "Alice Alpha", "2010-01-01"] + ["Correct"] * 60)
    ws_entry.append(["S2", "Bob Beta", "2010-02-02"] + ["Correct"] * 45 + ["Incorrect"] * 15)
    ws_entry.append(["S3", "Charlie Gamma", "2010-03-03"] + ["Correct"] * 30 + ["Unattempted"] * 30)
    ws_entry.append(["S4", "Diana Delta", "2010-04-04"] + ["Incorrect"] * 60)
    ws_entry.append(["S5", "Evan Epsilon", "2010-05-05"] + ["Correct"] * 50 + ["Incorrect"] * 5 + ["Unattempted"] * 5)
    
    wb.save(temp_file)
    
    result = validate_and_grade(temp_file)
    school_info = {"School Name": "Abako High", "Campus": "Main", "Grade": "10"}
    all_topics = ["Algebra", "Geometry", "Calculus"] # Calculus should be marked N/A
    
    generate_all_reports(result, school_info, all_topics)
    print("Reports generated!")
    os.remove(temp_file)

if __name__ == "__main__":
    test_reports()
